import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import kb_prep  # noqa: E402
from kb_prep import SHADOW_SUFFIX, main  # noqa: E402


def shadows(out: Path) -> list[str]:
    return sorted(str(p.relative_to(out)) for p in out.rglob("*" + SHADOW_SUFFIX))


def index(out: Path) -> list[dict]:
    """Wat kb_prep in de KB heeft staan, volgens het manifest."""
    files = json.loads((out / "_manifest.json").read_text(encoding="utf-8"))["files"]
    return [{"md": k, "source": v["source"]} for k, v in sorted(files.items())]


def test_shadow_name_keeps_extension_and_suffix(tmp_path):
    src, out = tmp_path / "kb", tmp_path / "out"
    (src / "sub").mkdir(parents=True)
    (src / "sub" / "notitie.txt").write_text("hallo", encoding="utf-8")
    (src / "tabel.csv").write_text("a;b\n1;2\n", encoding="utf-8")

    assert main([str(src), str(out)]) == 0
    assert shadows(out) == ["sub/notitie.txt" + SHADOW_SUFFIX, "tabel.csv" + SHADOW_SUFFIX]
    assert [e["md"] for e in index(out)] == ["sub/notitie.txt" + SHADOW_SUFFIX, "tabel.csv" + SHADOW_SUFFIX]


# 1. output binnen de bron wordt niet opnieuw ingelezen

@pytest.mark.parametrize("out_rel", ["_md", "."])
def test_output_inside_source_is_not_reingested(tmp_path, out_rel):
    src = tmp_path / "kb"
    src.mkdir()
    (src / "a.txt").write_text("inhoud", encoding="utf-8")
    (src / "eigen.md").write_text("# eigen notitie", encoding="utf-8")
    out = src / out_rel

    assert main([str(src), str(out)]) == 0
    first = shadows(out)
    assert main([str(src), str(out), "--force"]) == 0
    assert shadows(out) == first
    assert not any(SHADOW_SUFFIX + SHADOW_SUFFIX in s for s in first)
    assert not any("_index.json" in s for s in first)
    # de eigen .md van de gebruiker blijft staan en wordt gewoon omgezet
    assert (src / "eigen.md").exists()
    assert any(s.endswith("eigen.md" + SHADOW_SUFFIX) for s in first)


# 2. lege of niet-ondersteunde invoer crasht niet

def test_empty_source_writes_empty_manifest(tmp_path):
    src, out = tmp_path / "kb", tmp_path / "nog" / "niet" / "aanwezig"
    src.mkdir()
    assert main([str(src), str(out)]) == 0
    assert index(out) == []


def test_only_unsupported_files(tmp_path):
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    (src / "opname.mp3").write_bytes(b"ID3")
    assert main([str(src), str(out)]) == 0
    assert index(out) == []


# 3. verwijderde bronnen verdwijnen uit output en index

def test_deleted_source_is_removed(tmp_path):
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    (src / "blijft.txt").write_text("x", encoding="utf-8")
    (src / "weg.txt").write_text("y", encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert len(shadows(out)) == 2

    (src / "weg.txt").unlink()
    assert main([str(src), str(out)]) == 0
    assert shadows(out) == ["blijft.txt" + SHADOW_SUFFIX]
    assert [e["md"] for e in index(out)] == ["blijft.txt" + SHADOW_SUFFIX]


def test_cleanup_never_touches_other_md_files(tmp_path):
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    out.mkdir()
    (out / "handmatig.md").write_text("niet weggooien", encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert (out / "handmatig.md").exists()
    assert index(out) == []


# 4. mislukte OCR telt als mislukt, niet als tekst in de KB

def test_ocr_failure_is_reported_as_failure(tmp_path, monkeypatch):
    fitz = pytest.importorskip("pymupdf")
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    doc = fitz.open()
    doc.new_page()  # pagina zonder tekstlaag
    doc.save(str(src / "scan.pdf"))

    class Broken:
        @staticmethod
        def image_to_string(*_a, **_k):
            raise RuntimeError("tesseract niet gevonden")

    monkeypatch.setitem(sys.modules, "pytesseract", Broken)
    assert main([str(src), str(out), "--ocr"]) == 1
    assert shadows(out) == []
    assert index(out) == []


def test_ocr_failure_raises_from_converter(tmp_path, monkeypatch):
    fitz = pytest.importorskip("pymupdf")
    pdf = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf))
    monkeypatch.setitem(sys.modules, "pytesseract", None)  # import faalt
    with pytest.raises(RuntimeError, match="OCR mislukt op pagina 1"):
        kb_prep.conv_pdf(pdf, ocr=True)


def test_failed_reconversion_drops_old_shadow(tmp_path, monkeypatch):
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    (src / "data.json").write_text('{"a": 1}', encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert shadows(out) == ["data.json" + SHADOW_SUFFIX]

    def boom(path):
        raise ValueError("kapot")
    monkeypatch.setitem(kb_prep.CONVERTERS, ".json", boom)
    (src / "data.json").write_text('{"a": 2}', encoding="utf-8")
    assert main([str(src), str(out), "--force"]) == 1
    assert shadows(out) == []
    assert index(out) == []


# manifest en eigendom

def manifest(out: Path) -> dict:
    return json.loads((out / "_manifest.json").read_text(encoding="utf-8"))


def make_kb(tmp_path, files: dict[str, str]):
    src, out = tmp_path / "kb", tmp_path / "out"
    src.mkdir()
    for name, text in files.items():
        (src / name).parent.mkdir(parents=True, exist_ok=True)
        (src / name).write_text(text, encoding="utf-8")
    return src, out


def test_manifest_records_ownership(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src), str(out)]) == 0
    m = manifest(out)
    assert m["version"] == 1 and m["tool"] == "kb_prep"
    entry = m["files"]["a.txt" + SHADOW_SUFFIX]
    assert entry["source"] == str((src / "a.txt").resolve())
    assert entry["source_sha1"] == kb_prep.sha1_of(src / "a.txt")
    assert entry["shadow_sha1"] == kb_prep.sha1_of(out / ("a.txt" + SHADOW_SUFFIX))
    assert entry["ocr"] is True  # OCR staat standaard aan


def test_touched_but_unchanged_source_is_skipped(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    later = (src / "a.txt").stat().st_mtime + 100
    os.utime(src / "a.txt", (later, later))
    capsys.readouterr()
    assert main([str(src), str(out)]) == 0
    assert "0 omgezet, 1 overgeslagen" in capsys.readouterr().out


def test_changed_content_with_old_mtime_is_reconverted(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "oud"})
    main([str(src), str(out)])
    st = (src / "a.txt").stat()
    (src / "a.txt").write_text("nieuw", encoding="utf-8")
    os.utime(src / "a.txt", (st.st_atime, st.st_mtime))
    assert main([str(src), str(out)]) == 0
    assert "nieuw" in (out / ("a.txt" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_ocr_setting_change_only_reconverts_pdfs(tmp_path, capsys):
    fitz = pytest.importorskip("pymupdf")
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "tekstlaag")
    doc.save(str(src / "b.pdf"))
    main([str(src), str(out)])
    capsys.readouterr()
    assert main([str(src), str(out), "--no-ocr"]) == 0
    assert "1 omgezet, 1 overgeslagen" in capsys.readouterr().out
    assert manifest(out)["files"]["b.pdf" + SHADOW_SUFFIX]["ocr"] is False


def test_foreign_file_at_shadow_path_is_not_overwritten(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    out.mkdir()
    foreign = out / ("a.txt" + SHADOW_SUFFIX)
    foreign.write_text("---\nsource: \"ergens anders\"\n---\nvan hint-meet", encoding="utf-8")
    assert main([str(src), str(out), "--force"]) == 1
    assert "van hint-meet" in foreign.read_text(encoding="utf-8")
    assert manifest(out)["files"] == {}
    assert index(out) == []


def test_foreign_suffix_file_elsewhere_is_left_alone(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    out.mkdir()
    (out / ("notities-meeting.kb-hint-meet.md")).write_text("van hint-meet", encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert (out / "notities-meeting.kb-hint-meet.md").exists()
    assert [e["md"] for e in index(out)] == ["a.txt" + SHADOW_SUFFIX]


def test_edited_shadow_needs_force(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    shadow = out / ("a.txt" + SHADOW_SUFFIX)
    shadow.write_text(shadow.read_text(encoding="utf-8") + "\naantekening", encoding="utf-8")
    (src / "a.txt").write_text("y", encoding="utf-8")

    assert main([str(src), str(out)]) == 1
    assert "aantekening" in shadow.read_text(encoding="utf-8")
    assert main([str(src), str(out), "--force"]) == 0
    assert "aantekening" not in shadow.read_text(encoding="utf-8")


def test_edited_shadow_survives_source_deletion(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    shadow = out / ("a.txt" + SHADOW_SUFFIX)
    shadow.write_text("aangepast", encoding="utf-8")
    (src / "a.txt").unlink()
    assert main([str(src), str(out)]) == 1
    assert shadow.exists()
    assert manifest(out)["files"] == {}
    assert index(out) == []


def test_shadow_without_manifest_entry_is_never_overwritten(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    shadow = out / ("a.txt" + SHADOW_SUFFIX)
    shadow.write_text(shadow.read_text(encoding="utf-8") + "\neigen aantekening", encoding="utf-8")
    (out / "_manifest.json").unlink()
    assert main([str(src), str(out), "--force"]) == 1
    assert "eigen aantekening" in shadow.read_text(encoding="utf-8")


def test_manifest_cannot_delete_outside_output(tmp_path):
    src, out = make_kb(tmp_path, {})
    out.mkdir()
    victim = tmp_path / ("buiten" + SHADOW_SUFFIX)
    victim.write_text("niet aankomen", encoding="utf-8")
    sha = kb_prep.sha1_of(victim)
    bad = {"version": 1, "tool": "kb_prep", "files": {
        "../buiten" + SHADOW_SUFFIX: {"source": "x", "source_sha1": "x", "shadow_sha1": sha, "ocr": False},
        str(victim): {"source": "x", "source_sha1": "x", "shadow_sha1": sha, "ocr": False},
    }}
    (out / "_manifest.json").write_text(json.dumps(bad), encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert victim.exists()
    assert manifest(out)["files"] == {}


def test_foreign_manifest_aborts(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    out.mkdir()
    (out / "_manifest.json").write_text(json.dumps({"version": 1, "tool": "iets-anders", "files": {}}))
    with pytest.raises(SystemExit):
        main([str(src), str(out)])


def test_other_source_root_with_same_content_is_reconverted(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    src2 = tmp_path / "kb2"
    src2.mkdir()
    (src2 / "a.txt").write_text("x", encoding="utf-8")
    assert main([str(src2), str(out), "--force"]) == 0
    assert index(out)[0]["source"] == str((src2 / "a.txt").resolve())


def test_concurrent_run_is_refused(tmp_path):
    import fcntl
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    out.mkdir()
    with (out / ".kb_prep.lock").open("w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        assert main([str(src), str(out)]) == 3
    assert main([str(src), str(out)]) == 0


def test_manifest_saved_after_each_conversion(tmp_path, monkeypatch):
    src, out = make_kb(tmp_path, {"a.txt": "x", "b.txt": "y"})
    real = kb_prep.convert_one

    def crash_on_b(s, d, ocr, sha, **kw):
        if s.name == "b.txt":
            raise KeyboardInterrupt
        return real(s, d, ocr, sha, **kw)

    monkeypatch.setattr(kb_prep, "convert_one", crash_on_b)
    assert main([str(src), str(out)]) == 130
    assert list(manifest(out)["files"]) == ["a.txt" + SHADOW_SUFFIX]


@pytest.mark.skipif(os.geteuid() == 0, reason="root kan alles lezen")
def test_unreadable_source_is_a_failure_not_a_crash(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x", "b.txt": "y"})
    (src / "b.txt").chmod(0)
    try:
        assert main([str(src), str(out)]) == 1
    finally:
        (src / "b.txt").chmod(0o644)
    assert [e["md"] for e in index(out)] == ["a.txt" + SHADOW_SUFFIX]


def test_no_temp_files_left_behind(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x", "b.json": "{kapot"})
    main([str(src), str(out)])
    assert not list(out.rglob("*.tmp"))


# projecten

def test_project_writes_into_subfolder(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src), str(out), "--project", "fabrikam"]) == 0
    assert shadows(out / "fabrikam") == ["a.txt" + SHADOW_SUFFIX]
    assert (out / "fabrikam" / "_manifest.json").exists()
    assert not (out / "_manifest.json").exists()


def test_projects_are_independent(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    src2 = tmp_path / "kb2"
    src2.mkdir()
    (src2 / "b.txt").write_text("y", encoding="utf-8")
    main([str(src), str(out), "--project", "een"])
    main([str(src2), str(out), "--project", "twee"])
    before = {f: (out / "een" / f).read_bytes() for f in ("_manifest.json",)}

    # bron van project twee leeg: alles in twee wordt opgeruimd, een blijft ongemoeid
    (src2 / "b.txt").unlink()
    assert main([str(src2), str(out), "--project", "twee"]) == 0
    assert shadows(out / "twee") == []
    assert index(out / "twee") == []
    assert shadows(out / "een") == ["a.txt" + SHADOW_SUFFIX]
    assert {f: (out / "een" / f).read_bytes() for f in before} == before


@pytest.mark.parametrize("name", ["../weg", "a/b", ".verborgen", "_index", "..", ""])
def test_invalid_project_name_is_refused(tmp_path, name):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src), str(out), "--project", name]) == 2
    assert not out.exists()


# standaardwaarden

def test_ocr_is_default_and_no_ocr_skips_it(tmp_path, monkeypatch):
    fitz = pytest.importorskip("pymupdf")
    src, out = make_kb(tmp_path, {})
    doc = fitz.open()
    doc.new_page()  # scan zonder tekstlaag
    doc.save(str(src / "scan.pdf"))
    calls = []

    class Fake:
        @staticmethod
        def image_to_string(*_a, **_k):
            calls.append(1)
            return "gescande tekst"

    monkeypatch.setitem(sys.modules, "pytesseract", Fake)
    assert main([str(src), str(out / "met")]) == 0
    assert calls and "gescande tekst" in (out / "met" / ("scan.pdf" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    calls.clear()
    assert main([str(src), str(out / "zonder"), "--no-ocr"]) == 0
    assert not calls


@pytest.fixture
def kb_root(tmp_path, monkeypatch):
    root = tmp_path / "KB_md"
    monkeypatch.setenv("KB_ROOT", str(root))
    return root


def test_out_defaults_to_kb_root_and_source_name(tmp_path, kb_root):
    src, _ = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src)]) == 0
    assert shadows(kb_root / "kb") == ["a.txt" + SHADOW_SUFFIX]


def test_project_overrides_source_name(tmp_path, kb_root):
    src, _ = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src), "--project", "fabrikam"]) == 0
    assert shadows(kb_root / "fabrikam") == ["a.txt" + SHADOW_SUFFIX]


def test_source_name_that_is_no_valid_project_needs_flag(tmp_path, kb_root):
    src = tmp_path / "kado's"
    src.mkdir()
    assert main([str(src)]) == 2
    assert main([str(src), "--project", "kados"]) == 0


def test_kb_root_from_dotenv(tmp_path, monkeypatch):
    monkeypatch.delenv("KB_ROOT", raising=False)
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(f"KB_ROOT={tmp_path / 'uit_env'}\n", encoding="utf-8")
    src, _ = make_kb(tmp_path, {"a.txt": "x"})
    assert main([str(src)]) == 0
    assert shadows(tmp_path / "uit_env" / "kb") == ["a.txt" + SHADOW_SUFFIX]
    assert "KB_ROOT" not in os.environ  # .env lezen verandert de omgeving niet


def test_kb_root_fallback(monkeypatch, tmp_path):
    monkeypatch.delenv("KB_ROOT", raising=False)
    monkeypatch.chdir(tmp_path)
    assert kb_prep.default_root() == Path.home() / "KB_md"


def test_same_basename_from_other_folder_is_refused(tmp_path, kb_root):
    a = tmp_path / "klantA" / "docs"
    b = tmp_path / "klantB" / "docs"
    a.mkdir(parents=True)
    b.mkdir(parents=True)
    (a / "a.txt").write_text("van A", encoding="utf-8")
    (b / "b.txt").write_text("van B", encoding="utf-8")
    assert main([str(a)]) == 0
    assert main([str(b)]) == 2  # zou anders de schaduw van A opruimen
    assert shadows(kb_root / "docs") == ["a.txt" + SHADOW_SUFFIX]
    assert main([str(b), "--project", "docs-B"]) == 0


def test_manifest_records_source_root(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    main([str(src), str(out)])
    assert manifest(out)["source_root"] == str(src.resolve())


# OCR op slides en afbeeldingen, waarschuwing bij weinig tekst

class FakeOCR:
    calls: list = []

    @staticmethod
    def image_to_string(img, **_k):
        FakeOCR.calls.append(img.size)
        return "Belscript: eerst de klacht samenvatten, dan de vraag stellen."


@pytest.fixture
def fake_ocr(monkeypatch):
    FakeOCR.calls = []
    monkeypatch.setitem(sys.modules, "pytesseract", FakeOCR)
    return FakeOCR


def png(path: Path, size=(400, 200)):
    from PIL import Image
    Image.new("RGB", size, "white").save(path)


def picture_only_deck(path: Path, tmp_path: Path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    img = tmp_path / "slide.png"
    png(img)
    prs = pptx.Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # leeg layout
        slide.shapes.add_picture(str(img), Inches(0), Inches(0))
    prs.save(str(path))


def test_picture_only_slides_are_ocred(tmp_path, fake_ocr):
    src, out = make_kb(tmp_path, {})
    picture_only_deck(src / "deck.pptx", tmp_path)
    assert main([str(src), str(out)]) == 0
    text = (out / ("deck.pptx" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert text.count("Belscript") == 2
    assert len(fake_ocr.calls) == 2


def test_slides_with_text_are_not_ocred(tmp_path, fake_ocr):
    pptx = pytest.importorskip("pptx")
    src, out = make_kb(tmp_path, {})
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Titel"
    slide.placeholders[1].text = "Er staat al genoeg tekst op deze slide om niet te hoeven OCR-en."
    prs.save(str(src / "deck.pptx"))
    assert main([str(src), str(out)]) == 0
    assert fake_ocr.calls == []


def test_image_source_is_ocred(tmp_path, fake_ocr):
    src, out = make_kb(tmp_path, {})
    png(src / "belscript.png")
    assert main([str(src), str(out)]) == 0
    assert "Belscript" in (out / ("belscript.png" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_images_are_unsupported_without_ocr(tmp_path, fake_ocr, capsys):
    src, out = make_kb(tmp_path, {})
    png(src / "belscript.png")
    assert main([str(src), str(out), "--no-ocr"]) == 0
    assert shadows(out) == []
    assert "1 niet-ondersteund" in capsys.readouterr().out
    assert fake_ocr.calls == []


def test_switching_ocr_reconverts_images_and_decks(tmp_path, fake_ocr, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "genoeg tekst om niet als leeg gemeld te worden, echt waar."})
    picture_only_deck(src / "deck.pptx", tmp_path)
    main([str(src), str(out), "--no-ocr"])
    capsys.readouterr()
    assert main([str(src), str(out)]) == 0
    assert "1 omgezet, 1 overgeslagen" in capsys.readouterr().out


def test_low_text_is_warned_but_kept(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"leeg.txt": "", "vol.txt": "x" * 100})
    assert main([str(src), str(out)]) == 0
    printed = capsys.readouterr().out
    assert "⚠ leeg.txt: maar 0 tekens tekst" in printed
    assert "1 met weinig tekst" in printed
    assert "vol.txt" not in printed.split("Weinig tekst")[1]
    assert len(shadows(out)) == 2


def test_empty_deck_without_ocr_suggests_ocr(tmp_path, capsys):
    src, out = make_kb(tmp_path, {})
    picture_only_deck(src / "deck.pptx", tmp_path)
    main([str(src), str(out), "--no-ocr"])
    assert "deck.pptx: maar 0 tekens tekst, probeer zonder --no-ocr" in capsys.readouterr().out


def test_real_tesseract_reads_dutch(tmp_path):
    import shutil
    if not shutil.which("tesseract"):
        pytest.skip("tesseract niet geïnstalleerd")
    pytest.importorskip("pytesseract")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 120), "white")
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 48)
    except OSError:
        pytest.skip("geen systeemfont")
    ImageDraw.Draw(img).text((20, 30), "Geschillencommissie", fill="black", font=font)
    assert "Geschillencommissie" in kb_prep.ocr_image(img, "test")


def test_transparent_image_is_flattened_on_white(tmp_path):
    import shutil
    if not shutil.which("tesseract"):
        pytest.skip("tesseract niet geïnstalleerd")
    pytest.importorskip("pytesseract")
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 48)
    except OSError:
        pytest.skip("geen systeemfont")
    img = Image.new("RGBA", (900, 120), (0, 0, 0, 0))  # volledig transparant
    ImageDraw.Draw(img).text((20, 30), "Geschillencommissie", fill=(0, 0, 0, 255), font=font)
    assert "Geschillencommissie" in kb_prep.ocr_image(img, "test")


def test_new_converter_version_reconverts(tmp_path, monkeypatch, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100})
    main([str(src), str(out)])
    assert manifest(out)["files"]["a.txt" + SHADOW_SUFFIX]["converter"] == kb_prep.CONVERTER_VERSION
    monkeypatch.setattr(kb_prep, "CONVERTER_VERSION", kb_prep.CONVERTER_VERSION + 1)
    capsys.readouterr()
    main([str(src), str(out)])
    assert "1 omgezet" in capsys.readouterr().out


def test_multipage_tiff_reads_every_page(tmp_path, fake_ocr):
    from PIL import Image
    src, out = make_kb(tmp_path, {})
    pages = [Image.new("RGB", (100 + i, 100), "white") for i in range(3)]
    pages[0].save(src / "scan.tif", save_all=True, append_images=pages[1:])
    assert main([str(src), str(out)]) == 0
    text = (out / ("scan.tif" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert [f"## Pagina {i}" in text for i in (1, 2, 3)] == [True] * 3
    assert len(fake_ocr.calls) == 3


def test_grouped_picture_and_titled_scan_are_ocred(tmp_path, fake_ocr):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    src, out = make_kb(tmp_path, {})
    img = tmp_path / "scan.png"
    png(img)
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # alleen titel
    slide.shapes.title.text = "Bijlage"
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(str(img), Inches(1), Inches(1))
    prs.save(str(src / "deck.pptx"))
    assert main([str(src), str(out)]) == 0
    text = (out / ("deck.pptx" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "## Slide 1: Bijlage" in text and "Belscript" in text


def test_low_text_keeps_being_reported_when_skipped(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"leeg.txt": ""})
    main([str(src), str(out)])
    capsys.readouterr()
    main([str(src), str(out)])
    printed = capsys.readouterr().out
    assert "1 overgeslagen" in printed and "1 met weinig tekst" in printed


def test_ocr_uses_timeout(tmp_path, monkeypatch):
    from PIL import Image
    seen = {}

    class Fake:
        @staticmethod
        def image_to_string(img, **kw):
            seen.update(kw)
            return ""

    monkeypatch.setitem(sys.modules, "pytesseract", Fake)
    kb_prep.ocr_image(Image.new("RGB", (10, 10)), "test")
    assert seen["timeout"] == kb_prep.OCR_TIMEOUT


def test_no_index_file_and_legacy_index_removed(tmp_path):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    out.mkdir()
    (out / "_index.json").write_text("[]", encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert not (out / "_index.json").exists()


# voortgang

class FakeTTY:
    def __init__(self):
        self.buf = []

    def isatty(self):
        return True

    def write(self, s):
        self.buf.append(s)

    def flush(self):
        pass

    @property
    def text(self):
        return "".join(self.buf)


def test_progress_overwrites_one_line_in_terminal(tmp_path, monkeypatch, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100, "b.txt": "y" * 100, "leeg.txt": ""})
    tty = FakeTTY()
    monkeypatch.setattr(sys, "stderr", tty)
    monkeypatch.setattr(sys.stdout, "isatty", lambda: True)
    assert main([str(src), str(out)]) == 0
    printed = capsys.readouterr().out
    assert "[1/3] a.txt" in tty.text and "[3/3] leeg.txt" in tty.text
    assert "\r\033[K" in tty.text and "\n" not in tty.text  # alles op één regel
    assert "✓" not in printed  # routine zit in de voortgangsregel
    assert "⚠ leeg.txt" in printed  # meldingen blijven staan
    assert tty.buf[-1] == "\r\033[K"  # regel leeg voor de samenvatting


def test_no_progress_line_outside_terminal(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100})
    main([str(src), str(out)])
    captured = capsys.readouterr()
    assert "✓ a.txt" in captured.out
    assert "\r" not in captured.out + captured.err


def test_redirected_stdout_keeps_success_lines(tmp_path, monkeypatch, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100})
    tty = FakeTTY()
    monkeypatch.setattr(sys, "stderr", tty)  # stderr naar terminal, stdout naar bestand
    main([str(src), str(out)])
    assert "✓ a.txt" in capsys.readouterr().out
    assert tty.text == ""


def test_progress_line_cleared_on_interrupt(tmp_path, monkeypatch):
    src, out = make_kb(tmp_path, {"a.txt": "x"})
    tty = FakeTTY()
    monkeypatch.setattr(sys, "stderr", tty)
    monkeypatch.setattr(sys.stdout, "isatty", lambda: True)

    def boom(*_a, **_k):
        raise KeyboardInterrupt

    monkeypatch.setattr(kb_prep, "convert_one", boom)
    assert main([str(src), str(out)]) == 130
    assert "\r\033[K\nAfgebroken" in tty.text  # regel gewist vóór de melding


def test_page_progress_and_estimate(tmp_path, monkeypatch, capsys):
    fitz = pytest.importorskip("pymupdf")
    src, out = make_kb(tmp_path, {})
    doc = fitz.open()
    for i in range(4):
        doc.new_page().insert_text((72, 72), f"pagina {i} met genoeg tekst om niet als leeg gemeld te worden")
    doc.save(str(src / "groot.pdf"))
    clock = iter(range(0, 1000, 10))  # elke aanroep 10 s later
    monkeypatch.setattr(kb_prep.time, "monotonic", lambda: next(clock))
    tty = FakeTTY()
    monkeypatch.setattr(sys, "stderr", tty)
    monkeypatch.setattr(sys.stdout, "isatty", lambda: True)
    assert main([str(src), str(out)]) == 0
    frames = tty.text.split("\r\033[K")
    assert "[1/1] groot.pdf · pagina 1/4" in frames
    assert any(f.startswith("[1/1] groot.pdf · pagina 3/4 · nog ~") for f in frames)


def test_long_name_is_shortened_but_page_counter_kept(monkeypatch):
    tty = FakeTTY()
    monkeypatch.setattr(sys, "stderr", tty)
    monkeypatch.setattr(sys.stdout, "isatty", lambda: True)
    monkeypatch.setattr(kb_prep.shutil, "get_terminal_size", lambda *_: os.terminal_size((50, 20)))
    p = kb_prep.Progress(375)
    p.update(7, "01-woning-en-installatie/AFSCHRIFT akte Splitsing blok C Florence.pdf")
    p.tick(22, 80, "pagina")
    last = tty.text.split("\r\033[K")[-1]
    assert len(last) <= 49 and last.endswith("pagina 23/80") and "…" in last


def test_fmt_duration():
    assert kb_prep.fmt_duration(40) == "40 s"
    assert kb_prep.fmt_duration(300) == "5 min"
    assert kb_prep.fmt_duration(3 * 3600) == "3.0 u"



def test_ctrl_c_gives_short_message(tmp_path, monkeypatch, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x"})

    def boom(*_a, **_k):
        raise KeyboardInterrupt

    monkeypatch.setattr(kb_prep, "convert_one", boom)
    assert main([str(src), str(out)]) == 130
    err = capsys.readouterr().err
    assert "Afgebroken" in err and "Traceback" not in err


# docx

def test_docx_paragraph_without_style(tmp_path):
    docx = pytest.importorskip("docx")
    src, out = make_kb(tmp_path, {})
    d = docx.Document()
    d.add_paragraph("Paragraaf met stijl")
    p = d.add_paragraph("Paragraaf zonder stijl")  # geen w:pStyle én (hieronder) geen standaardstijl
    styles = d.styles.element
    for st in styles.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}style"):
        if st.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}default") == "1":
            styles.remove(st)
    assert p.style is None  # zoals in de Contoso-memo
    d.save(str(src / "memo.docx"))
    assert main([str(src), str(out)]) == 0
    assert "Paragraaf zonder stijl" in (out / ("memo.docx" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_docx_with_only_screenshots_is_ocred(tmp_path, fake_ocr):
    docx = pytest.importorskip("docx")
    src, out = make_kb(tmp_path, {})
    img = tmp_path / "screenshot.png"
    png(img)
    d = docx.Document()
    d.add_picture(str(img))
    d.add_picture(str(img))
    d.save(str(src / "logging.docx"))
    assert main([str(src), str(out)]) == 0
    assert "Belscript" in (out / ("logging.docx" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_docx_with_text_is_not_ocred(tmp_path, fake_ocr):
    docx = pytest.importorskip("docx")
    src, out = make_kb(tmp_path, {})
    img = tmp_path / "grafiek.png"
    png(img)
    d = docx.Document()
    d.add_paragraph("Ruim voldoende tekst in dit document, dus de grafiek hoeft niet door OCR.")
    d.add_picture(str(img))
    d.save(str(src / "rapport.docx"))
    main([str(src), str(out)])
    assert fake_ocr.calls == []


# foto's

def test_photo_without_text_is_kept_quietly(tmp_path, monkeypatch, capsys):
    class Noise:
        @staticmethod
        def image_to_string(*_a, **_k):
            return "~ |"

    monkeypatch.setitem(sys.modules, "pytesseract", Noise)
    src, out = make_kb(tmp_path, {})
    png(src / "20260625Waterdrukmeter.jpg")
    assert main([str(src), str(out)]) == 0
    printed = capsys.readouterr().out
    assert "⚠" not in printed and "1 foto's met weinig of geen tekst" in printed
    shadow = (out / ("20260625Waterdrukmeter.jpg" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "# 20260625Waterdrukmeter" in shadow and kb_prep.PHOTO_NOTE in shadow
    assert "Herkende tekst (OCR, onzeker): ~ |" in shadow  # kort, maar bewaard
    capsys.readouterr()
    main([str(src), str(out)])  # overgeslagen: nog steeds als foto geteld, niet als waarschuwing
    printed = capsys.readouterr().out
    assert "1 overgeslagen" in printed and "1 foto's met weinig of geen tekst" in printed and "⚠" not in printed


def test_type_version_only_redoes_that_type(tmp_path, monkeypatch, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100, "b.csv": "a,b\n" + "1,2\n" * 30})
    main([str(src), str(out)])
    monkeypatch.setitem(kb_prep.TYPE_VERSION, ".csv", 99)
    capsys.readouterr()
    main([str(src), str(out)])
    assert "1 omgezet, 1 overgeslagen" in capsys.readouterr().out


def test_photo_with_short_real_text_keeps_it(tmp_path, monkeypatch, capsys):
    class Pin:
        @staticmethod
        def image_to_string(*_a, **_k):
            return "Enter\nPIN"

    monkeypatch.setitem(sys.modules, "pytesseract", Pin)
    src, out = make_kb(tmp_path, {})
    png(src / "pincode.jpeg")
    main([str(src), str(out)])
    assert "⚠" not in capsys.readouterr().out
    shadow = (out / ("pincode.jpeg" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "Herkende tekst (OCR, onzeker): Enter PIN" in shadow


def test_blank_photo_has_note_only(tmp_path, monkeypatch):
    class Blank:
        @staticmethod
        def image_to_string(*_a, **_k):
            return ""

    monkeypatch.setitem(sys.modules, "pytesseract", Blank)
    src, out = make_kb(tmp_path, {})
    png(src / "vloer.jpg")
    main([str(src), str(out)])
    shadow = (out / ("vloer.jpg" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert shadow.rstrip().endswith(kb_prep.PHOTO_NOTE)


# e-mail

def write_eml(path: Path, *, html=False, attachments=(), forwarded=None):
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["From"] = "Contoso <particulieren@contoso.nl>"
    msg["To"] = "stefan@example.com"
    msg["Date"] = "Fri, 20 Feb 2026 11:36:53 +0000"
    msg["Subject"] = "Bevestiging onderhoudsabonnement"
    if html:
        msg.set_content("<html><body><h1>Offerte</h1><p>Wij bevestigen <b>S300713739</b>.</p>"
                        "<img src='cid:logo'></body></html>", subtype="html")
    else:
        msg.set_content("Beste Stefan,\n\nHierbij de bevestiging van uw abonnement.")
    for name, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    if forwarded is not None:
        msg.add_attachment(forwarded)
    path.write_bytes(bytes(msg))


def test_eml_headers_and_plain_body(tmp_path):
    src, out = make_kb(tmp_path, {})
    write_eml(src / "bevestiging.eml")
    assert main([str(src), str(out)]) == 0
    text = (out / ("bevestiging.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "- **Van:** Contoso <particulieren@contoso.nl>" in text
    assert "- **Onderwerp:** Bevestiging onderhoudsabonnement" in text
    assert "bevestiging van uw abonnement" in text


def test_eml_html_body_becomes_markdown(tmp_path):
    src, out = make_kb(tmp_path, {})
    write_eml(src / "offerte.eml", html=True)
    main([str(src), str(out)])
    text = (out / ("offerte.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "# Offerte" in text and "**S300713739**" in text and "<p>" not in text


def test_eml_attachments_are_converted_or_listed(tmp_path):
    fitz = pytest.importorskip("pymupdf")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Algemene voorwaarden onderhoudsabonnement")
    pdf = doc.tobytes()
    src, out = make_kb(tmp_path, {})
    write_eml(src / "mail.eml", attachments=[
        ("voorwaarden.pdf", pdf, "application", "pdf"),
        ("notitie.txt", "Kenmerk 3-225596".encode(), "text", "plain"),
        ("archief.zip", b"PK", "application", "zip"),
        ("kapot.pdf", b"geen pdf", "application", "pdf"),
    ])
    assert main([str(src), str(out)]) == 0  # kapotte bijlage laat de mail niet mislukken
    text = (out / ("mail.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "## Bijlage: voorwaarden.pdf" in text and "Algemene voorwaarden" in text
    assert "## Bijlage: notitie.txt" in text and "3-225596" in text
    assert "## Bijlage: archief.zip\n\n_(niet gelezen: bestandstype)_" in text
    assert "## Bijlage: kapot.pdf\n\n_(niet gelezen:" in text


def test_eml_forwarded_message(tmp_path):
    from email.message import EmailMessage
    inner = EmailMessage()
    inner["From"] = "werkvoorbereiding@contoso.nl"
    inner["Subject"] = "Offerte S300713739"
    inner.set_content("De oorspronkelijke offerte.")
    src, out = make_kb(tmp_path, {})
    write_eml(src / "fw.eml", forwarded=inner)
    main([str(src), str(out)])
    text = (out / ("fw.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "## Doorgestuurd bericht" in text and "De oorspronkelijke offerte." in text


# .kbignore

def test_kbignore_excludes_and_cleans_up(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100, "debug.log.txt": "y" * 100})
    (src / "tmp").mkdir()
    (src / "tmp" / "render.txt").write_text("z" * 100, encoding="utf-8")
    (src / "tmp" / "bewaar.txt").write_text("w" * 100, encoding="utf-8")
    main([str(src), str(out)])
    assert len(shadows(out)) == 4

    (src / ".kbignore").write_text("# tussenbestanden\ntmp/\n!tmp/bewaar.txt\n*.log.txt\n", encoding="utf-8")
    capsys.readouterr()
    assert main([str(src), str(out)]) == 0
    printed = capsys.readouterr().out
    # ! werkt ook binnen een uitgesloten map (vriendelijker dan git zelf)
    assert shadows(out) == ["a.txt" + SHADOW_SUFFIX, "tmp/bewaar.txt" + SHADOW_SUFFIX]
    assert "2 opgeruimd" in printed and "2 genegeerd via .kbignore" in printed


def test_kbignore_negation_inside_file_pattern(tmp_path):
    src, out = make_kb(tmp_path, {"a.log.txt": "x" * 100, "keep.log.txt": "y" * 100})
    (src / ".kbignore").write_text("*.log.txt\n!keep.log.txt\n", encoding="utf-8")
    main([str(src), str(out)])
    assert shadows(out) == ["keep.log.txt" + SHADOW_SUFFIX]


def test_no_kbignore_means_nothing_ignored(tmp_path, capsys):
    src, out = make_kb(tmp_path, {"a.txt": "x" * 100})
    main([str(src), str(out)])
    assert "genegeerd" not in capsys.readouterr().out


def test_eml_inline_images_are_skipped(tmp_path):
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["Subject"] = "RE: Offerte"
    msg.set_content("Zie offerte.")
    msg.add_alternative("<p>Zie offerte.</p><img src='cid:logo'>", subtype="html")
    msg.get_payload()[1].add_related(b"GIF89a", maintype="image", subtype="gif", cid="<logo>",
                                     filename="image001.gif")
    msg.add_attachment(b"%PDF-kapot", maintype="application", subtype="pdf", filename="offerte.pdf")
    src, out = make_kb(tmp_path, {})
    (src / "re.eml").write_bytes(bytes(msg))
    main([str(src), str(out)])
    text = (out / ("re.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert "image001.gif" not in text and "## Bijlage: offerte.pdf" in text


def test_eml_text_attachment_keeps_declared_charset(tmp_path):
    src, out = make_kb(tmp_path, {})
    write_eml(src / "mail.eml", attachments=[("notitie.txt", "Café über €".encode("cp1252"), "text", "plain")])
    # zet de bijlage om naar windows-1252 met charset in de kop, zoals Outlook dat doet
    from email import policy
    from email.parser import BytesParser
    msg = BytesParser(policy=policy.default).parsebytes((src / "mail.eml").read_bytes())
    att = next(msg.iter_attachments())
    att.set_content("Café über €", subtype="plain", charset="windows-1252", disposition="attachment",
                    filename="notitie.txt")
    (src / "mail.eml").write_bytes(bytes(msg))
    main([str(src), str(out)])
    assert "Café über €" in (out / ("mail.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_eml_broken_forward_does_not_fail_mail(tmp_path, monkeypatch):
    from email.message import EmailMessage
    inner = EmailMessage()
    inner["Subject"] = "binnen"
    inner.set_content("x")
    src, out = make_kb(tmp_path, {})
    write_eml(src / "fw.eml", forwarded=inner)
    real = kb_prep.render_email

    def flaky(msg, ocr, tick, depth=0):
        if depth == 1:
            raise ValueError("kapot")
        return real(msg, ocr, tick, depth)

    monkeypatch.setattr(kb_prep, "render_email", flaky)
    assert main([str(src), str(out)]) == 0
    assert "_(niet gelezen: ValueError: kapot)_" in (out / ("fw.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_eml_nesting_is_capped(tmp_path):
    from email.message import EmailMessage
    msg = EmailMessage()
    msg["Subject"] = "niveau 0"
    msg.set_content("0")
    for level in range(1, 8):
        outer = EmailMessage()
        outer["Subject"] = f"niveau {level}"
        outer.set_content(str(level))
        outer.add_attachment(msg)
        msg = outer
    src, out = make_kb(tmp_path, {})
    (src / "diep.eml").write_bytes(bytes(msg))
    assert main([str(src), str(out)]) == 0
    assert "dieper dan 5 niveaus" in (out / ("diep.eml" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


@pytest.mark.parametrize("ext", [".pdf", ".docx"])  # .docx: ophoging onder het maximum
def test_eml_version_follows_other_converters(monkeypatch, ext):
    before = kb_prep.converter_version(".eml")
    monkeypatch.setitem(kb_prep.TYPE_VERSION, ext, kb_prep.converter_version(ext) + 1)
    assert kb_prep.converter_version(".eml") != before


def test_long_source_name_gets_short_unique_shadow(tmp_path):
    base = "akte_" + "é" * 115  # 235 bytes: past zelf, met suffix niet meer
    src, out = make_kb(tmp_path, {base + "a.txt": "eerste " * 20, base + "b.txt": "tweede " * 20})
    assert main([str(src), str(out)]) == 0
    names = shadows(out)
    assert len(names) == 2
    assert all(len(n.encode()) <= kb_prep.MAX_SHADOW_NAME and "~" in n for n in names)
    assert main([str(src), str(out)]) == 0  # tweede run: alles actueel, geen conflicten
    assert shadows(out) == names


def test_json_with_comments_is_kept_as_text(tmp_path):
    src, out = make_kb(tmp_path, {"tsconfig.json": '{\n  // commentaar\n  "strict": true, "target": "es2020", "module": "commonjs"\n}'})
    assert main([str(src), str(out)]) == 0
    assert "// commentaar" in (out / ("tsconfig.json" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_dependency_and_hidden_dirs_are_skipped(tmp_path):
    long = "inhoud die lang genoeg is om niet als weinig tekst te tellen"
    src, out = make_kb(tmp_path, {"a.txt": long, "node_modules/x/b.txt": long, ".git/c.txt": long})
    assert main([str(src), str(out)]) == 0
    assert shadows(out) == ["a.txt" + SHADOW_SUFFIX]


def test_zero_filled_file_fails_with_clear_message(tmp_path, capsys):
    src, out = make_kb(tmp_path, {})
    (src / "kopie.pdf").write_bytes(bytes(5000))
    assert main([str(src), str(out)]) == 1
    assert "alleen nullen" in capsys.readouterr().out


def test_docx_with_missing_content_type_is_repaired(tmp_path):
    import zipfile
    from docx import Document
    src, out = make_kb(tmp_path, {})
    good = tmp_path / "goed.docx"
    d = Document()
    d.add_paragraph("Overeenkomst van geldlening tussen partijen, met voldoende tekst erin.")
    d.save(good)
    with zipfile.ZipFile(good) as zin, zipfile.ZipFile(src / "lening.docx", "w") as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr("word/fonts/Arimo-regular.ttf", b"font")
    assert main([str(src), str(out)]) == 0
    assert "geldlening" in (out / ("lening.docx" + SHADOW_SUFFIX)).read_text(encoding="utf-8")


def test_tall_image_is_ocrd_in_strips(monkeypatch):
    from PIL import Image
    import pytesseract
    sizes = []
    monkeypatch.setattr(pytesseract, "image_to_string", lambda img, **kw: sizes.append(img.size) or "x")
    kb_prep.ocr_image(Image.new("L", (600, kb_prep.OCR_STRIP * 2 + 10)), "test")
    assert sizes == [(600, kb_prep.OCR_STRIP)] * 2 + [(600, 10)]


def test_build_output_is_skipped_only_next_to_a_project_file(tmp_path):
    long = "inhoud die lang genoeg is om niet als weinig tekst te tellen"
    src, out = make_kb(tmp_path, {
        "app/package.json": '{"name": "app", "version": "1.0.0", "description": "testproject"}',
        "app/dist/bundle.txt": long, "app/x.egg-info/PKG-INFO.txt": long,
        "admin/dist/factuur.txt": long,  # geen projectbestand ernaast: gewoon een map
    })
    assert main([str(src), str(out)]) == 0
    assert shadows(out) == ["admin/dist/factuur.txt" + SHADOW_SUFFIX, "app/package.json" + SHADOW_SUFFIX]


def test_encrypted_pdf_opens_with_password_from_kbpasswords(tmp_path):
    import pymupdf as fitz
    src, out = make_kb(tmp_path, {})
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Contract met voldoende tekst om mee te tellen in de KB.")
    doc.save(src / "contract.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="geheim", owner_pw="eigenaar")
    assert main([str(src), str(out)]) == 1  # zonder wachtwoord: mislukt

    (src / ".kbpasswords").write_text("# wachtwoorden\nverkeerd\ngeheim\n", encoding="utf-8")
    assert main([str(src), str(out)]) == 0
    assert "Contract met voldoende tekst" in (out / ("contract.pdf" + SHADOW_SUFFIX)).read_text(encoding="utf-8")
    assert not (out / (".kbpasswords" + SHADOW_SUFFIX)).exists()
